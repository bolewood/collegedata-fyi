"""Parse the FSA institutional nonpayment workbook."""

from __future__ import annotations

import hashlib
import re
from collections import Counter
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from tools.fsa.opeid import restore_fsa_opeid

SCHOOL_LEVEL_SHEET = "School-Level Nonpayment Rates"
EXPECTED_COLUMNS = (
    "OPE ID",
    "School Name",
    "School Type",
    "State",
    "Total Borrowers Evaluated",
    "Nonpayment Rate",
)
RATE_SUPPRESSION_TOKENS = ("<10%", "<5%", "<4%", "<3%", "<2%", "Not Calculated")
DENOM_SUPPRESSION_TOKENS = ("<100",)
SMALL_N_CUT = 10
SOURCE_URL = (
    "https://studentaid.gov/sites/default/files/fsawg/datacenter/library/"
    "nonpayment-rates.xlsx"
)
LISTING_PAGE = "https://studentaid.gov/data-center/student/portfolio"


class ColumnDriftError(ValueError):
    """Raised when the school-level sheet is missing expected headers."""


@dataclass(frozen=True)
class FsaRow:
    opeid: str
    school_name: str
    school_type: str | None
    state: str | None
    borrowers_in_denom: int | None
    borrowers_raw: str | None
    nonpayment_rate: float | None
    rate_raw: str | None
    suppressed: bool
    public_visible: bool


@dataclass(frozen=True)
class ParsedWorkbook:
    title: str
    as_of_label: str
    as_of_date: date
    cohort_window_start: date
    cohort_window_end: date
    data_run_note: str
    header_row: int
    opeid_digit_length: int
    opeid_storage: str
    rate_scale: str
    suppression_tokens: tuple[str, ...]
    rows: tuple[FsaRow, ...]
    skipped_blank: int


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _cell_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _parse_denom(value: Any) -> tuple[int | None, str | None, bool]:
    if value is None:
        return None, None, True
    if isinstance(value, bool):
        return None, str(value), True
    if isinstance(value, (int, float)) and value == value:
        if isinstance(value, float) and value != int(value):
            return None, str(value), True
        number = int(value)
        return number, str(number), number < SMALL_N_CUT
    text = str(value).strip()
    if text in DENOM_SUPPRESSION_TOKENS:
        return None, text, True
    if text.isdigit():
        number = int(text)
        return number, text, number < SMALL_N_CUT
    return None, text, True


def _parse_rate(value: Any) -> tuple[float | None, str | None, bool, str]:
    if value is None:
        return None, None, True, "missing"
    if isinstance(value, str):
        text = value.strip()
        if text in RATE_SUPPRESSION_TOKENS:
            return None, text, True, "token"
        try:
            number = float(text.replace("%", ""))
        except ValueError:
            return None, text, True, "token"
        value = number
        raw = text
    else:
        raw = str(value)
        number = float(value)
    if number > 1.0:
        if number > 100:
            return None, raw, True, "invalid"
        return number / 100.0, raw, False, "0-100"
    return number, raw, False, "0-1"


def _public_visible(rate: float | None, denom: int | None, suppressed: bool) -> bool:
    if suppressed or rate is None:
        return False
    if denom is not None and denom < SMALL_N_CUT:
        return False
    return 0.0 <= rate <= 1.0


def _header_row(ws: Worksheet) -> tuple[int, list[str]]:
    for index, row in enumerate(ws.iter_rows(min_row=1, max_row=8, values_only=True), 1):
        values = [_cell_text(cell) or "" for cell in row[:6]]
        if values[:6] == list(EXPECTED_COLUMNS):
            return index, values
    raise ColumnDriftError(
        f"{SCHOOL_LEVEL_SHEET} is missing expected headers {EXPECTED_COLUMNS}"
    )


def _as_of_from_title(title: str, definitions_note: str) -> tuple[str, date]:
    label_match = re.search(r"as of\s+([A-Za-z]+ \d{4})", title, re.I)
    run_match = re.search(r"late\s+([A-Za-z]+ \d{4})", definitions_note, re.I)
    label = None
    if run_match:
        label = f"late {run_match.group(1)}"
    elif label_match:
        label = label_match.group(1)
    if not label:
        raise ValueError("Could not read a public as-of date from the workbook")
    month_year = re.search(r"([A-Za-z]+)\s+(\d{4})", label)
    if not month_year:
        raise ValueError(f"Unparseable as-of label: {label}")
    months = {
        "january": 1, "february": 2, "march": 3, "april": 4, "may": 5, "june": 6,
        "july": 7, "august": 8, "september": 9, "october": 10, "november": 11,
        "december": 12,
    }
    month = months[month_year.group(1).lower()]
    year = int(month_year.group(2))
    # "late May 2026" has no calendar day. Use month-end as the dated pull.
    as_of = date(year, month, 31 if month == 5 else 28 if month == 2 else 30)
    if month in {1, 3, 7, 8, 10, 12}:
        as_of = date(year, month, 31)
    elif month in {4, 6, 9, 11}:
        as_of = date(year, month, 30)
    elif month == 2:
        as_of = date(year, month, 29 if year % 4 == 0 else 28)
    return label, as_of


def parse_workbook(path: Path) -> ParsedWorkbook:
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        if SCHOOL_LEVEL_SHEET not in wb.sheetnames:
            raise ColumnDriftError(
                f"Workbook is missing sheet {SCHOOL_LEVEL_SHEET!r}; "
                f"found {wb.sheetnames}"
            )
        definitions = " ".join(
            str(cell[0]) for cell in wb["Definitions"].iter_rows(values_only=True)
            if cell and cell[0]
        )
        ws = wb[SCHOOL_LEVEL_SHEET]
        title = _cell_text(next(ws.iter_rows(min_row=1, max_row=1, values_only=True))[0]) or ""
        header_row, headers = _header_row(ws)
        missing = [col for col in EXPECTED_COLUMNS if col not in headers]
        if missing:
            raise ColumnDriftError(f"Renamed or missing columns: {missing}")

        rows: list[FsaRow] = []
        skipped_blank = 0
        opeid_lens: Counter[int] = Counter()
        storage_kinds: Counter[str] = Counter()
        rate_scales: Counter[str] = Counter()
        tokens: set[str] = set()

        for raw in ws.iter_rows(min_row=header_row + 1, max_col=6, values_only=True):
            opeid_raw, name, school_type, state, denom, rate = (list(raw) + [None] * 6)[:6]
            if opeid_raw is None and name is None and rate is None:
                skipped_blank += 1
                continue
            opeid = restore_fsa_opeid(opeid_raw)
            if opeid is None:
                skipped_blank += 1
                continue
            storage_kinds["numeric" if isinstance(opeid_raw, (int, float)) else "text"] += 1
            opeid_lens[len(opeid)] += 1
            denom_n, denom_raw, denom_suppressed = _parse_denom(denom)
            rate_n, rate_raw, rate_suppressed, scale = _parse_rate(rate)
            rate_scales[scale] += 1
            if rate_raw in RATE_SUPPRESSION_TOKENS or denom_raw in DENOM_SUPPRESSION_TOKENS:
                tokens.add(rate_raw or denom_raw or "")
            suppressed = denom_suppressed or rate_suppressed
            rows.append(
                FsaRow(
                    opeid=opeid,
                    school_name=_cell_text(name) or "",
                    school_type=_cell_text(school_type),
                    state=_cell_text(state),
                    borrowers_in_denom=denom_n,
                    borrowers_raw=denom_raw,
                    nonpayment_rate=rate_n,
                    rate_raw=rate_raw,
                    suppressed=suppressed,
                    public_visible=_public_visible(rate_n, denom_n, suppressed),
                )
            )
    finally:
        wb.close()

    if not rows:
        raise ValueError("School-level sheet had no OPE ID rows")
    duplicate = _duplicate_opeids(rows)
    if duplicate:
        raise ValueError(f"Duplicate FSA OPEID(s) in one workbook: {sorted(duplicate)[:8]}")

    digit_length = max(opeid_lens, key=opeid_lens.get)
    storage = max(storage_kinds, key=storage_kinds.get)
    rate_scale = "0-1" if rate_scales.get("0-1", 0) >= rate_scales.get("0-100", 0) else "0-100"
    as_of_label, as_of_date = _as_of_from_title(title, definitions)
    return ParsedWorkbook(
        title=title,
        as_of_label=as_of_label,
        as_of_date=as_of_date,
        cohort_window_start=date(2020, 1, 1),
        cohort_window_end=date(2025, 5, 31),
        data_run_note=definitions,
        header_row=header_row,
        opeid_digit_length=digit_length,
        opeid_storage=storage,
        rate_scale=rate_scale,
        suppression_tokens=tuple(sorted(t for t in tokens if t)),
        rows=tuple(rows),
        skipped_blank=skipped_blank,
    )


def _duplicate_opeids(rows: Iterable[FsaRow]) -> set[str]:
    seen: set[str] = set()
    dupes: set[str] = set()
    for row in rows:
        if row.opeid in seen:
            dupes.add(row.opeid)
        seen.add(row.opeid)
    return dupes
